"""Typed loaders for the frozen official M6 files (see data/raw/m6_official/README.md).

The raw snapshot stores every CSV column as a string; typing happens here so the
raw files stay a lossless copy of upstream.
"""

from __future__ import annotations

import hashlib
import json
from functools import cache
from pathlib import Path

import pandas as pd

REPO_ROOT = Path(__file__).resolve().parents[3]
MANIFEST = REPO_ROOT / "data" / "manifest.json"
SNAPSHOT = "m6_official"

RANK_COLUMNS = ["Rank1", "Rank2", "Rank3", "Rank4", "Rank5"]
# Order of the evaluation months in period 1 ("Trial run" is the pilot).
EVALUATIONS = ["Trial run"] + [
    f"{n}{'st' if n == 1 else 'nd' if n == 2 else 'rd' if n == 3 else 'th'} Submission"
    for n in range(1, 13)
]


@cache
def manifest() -> dict:
    return json.loads(MANIFEST.read_text(encoding="utf-8"))


def snapshot_files(downloaded: str | None = None) -> dict[str, dict]:
    """Manifest entries of the M6 snapshot, keyed by stored name (e.g. ``"assets_m6"``).

    Uses the most recent snapshot unless ``downloaded`` (YYYY-MM-DD) is given.
    """
    snaps = [s for s in manifest()["snapshots"] if s["name"] == SNAPSHOT]
    if downloaded:
        snaps = [s for s in snaps if s["downloaded"] == downloaded]
    if not snaps:
        raise FileNotFoundError(f"no {SNAPSHOT} snapshot in {MANIFEST}")
    snap = max(snaps, key=lambda s: s["downloaded"])
    return {Path(f["path"]).stem.rsplit("_", 1)[0]: f for f in snap["files"]}


def path(name: str, downloaded: str | None = None) -> Path:
    return REPO_ROOT / snapshot_files(downloaded)[name]["path"]


def verify(downloaded: str | None = None) -> None:
    """Raise if any stored file does not match its manifest SHA-256."""
    for name, entry in snapshot_files(downloaded).items():
        digest = hashlib.sha256((REPO_ROOT / entry["path"]).read_bytes()).hexdigest()
        if digest != entry["sha256"]:
            raise ValueError(f"{name}: sha256 mismatch ({entry['path']})")


def load_assets(downloaded: str | None = None) -> pd.DataFrame:
    """Daily adjusted close prices: ``symbol`` (str), ``date`` (datetime64), ``price`` (float)."""
    df = pd.read_parquet(path("assets_m6", downloaded))
    return pd.DataFrame(
        {
            "symbol": df["symbol"].astype(str),
            "date": pd.to_datetime(df["date"], format="%Y/%m/%d"),
            "price": df["price"].astype(float),
        }
    )


def load_submissions(downloaded: str | None = None) -> pd.DataFrame:
    """Team submissions of period 1 with numeric columns typed.

    ``Evaluation`` is the month the submission was scored in; when a team did not
    resubmit, its previous submission is carried forward (``Submission`` differs).
    """
    df = pd.read_parquet(path("submissions", downloaded))
    out = df[["Team", "Submission", "Evaluation", "Symbol"]].astype(str).copy()
    for col in ["Decision", *RANK_COLUMNS]:
        out[col] = df[col].astype(float)
    out["IsActive"] = df["IsActive"].astype(int)
    for col in ["RpsHashCode", "IrHashCode"]:
        out[col] = pd.to_numeric(df[col].replace("", None)).astype("Int64")
    return out


def load_template(downloaded: str | None = None) -> pd.DataFrame:
    """Official submission template: uniform probabilities and equal weights (0.01).

    The template predates Facebook's ticker change and lists ``FB``; prices and
    submissions use ``META`` for the whole period, so it is renamed here.
    """
    df = pd.read_parquet(path("template", downloaded))
    out = df[["ID"]].astype(str).replace({"ID": {"FB": "META"}}).copy()
    for col in [*RANK_COLUMNS, "Decision"]:
        out[col] = df[col].astype(float)
    return out


def load_period1_calendar(downloaded: str | None = None) -> pd.DataFrame:
    """Evaluation windows of period 1, read from the official worked example.

    Each sheet of ``Evaluation - example.xlsx`` (Pilot, Month1..Month12) lists, in
    row 4 from column B, the base date (last close before the period) followed by the
    eligible days of the period. Returns one row per evaluation with ``evaluation``
    (as in ``submissions.Evaluation``), ``sheet``, ``base_date``, ``end_date`` and
    ``n_days`` (eligible days, base date excluded).
    """
    import openpyxl

    wb = openpyxl.load_workbook(path("evaluation_example", downloaded), read_only=True)
    sheets = ["Pilot"] + [f"Month{n}" for n in range(1, 13)]
    rows = []
    for evaluation, sheet in zip(EVALUATIONS, sheets, strict=True):
        header = next(wb[sheet].iter_rows(min_row=4, max_row=4, values_only=True))
        dates = []
        for value in header[1:]:
            if not hasattr(value, "year"):
                break
            dates.append(pd.Timestamp(value))
        rows.append(
            {
                "evaluation": evaluation,
                "sheet": sheet,
                "base_date": dates[0],
                "end_date": dates[-1],
                "n_days": len(dates) - 1,
            }
        )
    wb.close()
    return pd.DataFrame(rows)


def load_leaderboard(level: str = "month", downloaded: str | None = None) -> pd.DataFrame:
    """Official leaderboard from ``summary_leaderboard.xlsx``.

    ``level``: ``"month"`` (sheets "Trial" and "Month"; column ``evaluation`` as in
    ``submissions.Evaluation``), ``"quarter"`` (column ``period`` 1-4) or ``"global"``.
    Columns: ``team`` (8-character id, as in ``submissions.Team``), ``team_name``,
    ``rps`` and ``ir`` (rounded to 5 decimals upstream), ``rps_rank``, ``ir_rank``,
    ``overall_rank`` and ``position``.
    """
    import openpyxl

    sheets = {"month": ("Trial", "Month"), "quarter": ("Quarter",), "global": ("Global",)}
    if level not in sheets:
        raise ValueError(f"unknown level {level!r}")
    wb = openpyxl.load_workbook(path("summary_leaderboard", downloaded), read_only=True)
    rows = []
    for sheet in sheets[level]:
        # Every sheet but "Global" starts with a formula "Key" column.
        o = 0 if sheet == "Global" else 1
        for r in wb[sheet].iter_rows(min_row=2, values_only=True):
            if r[o + 1] is None:
                continue
            team_id, _, name = str(r[o + 1]).partition("\xa0")
            row = {
                "team": team_id,
                "team_name": name,
                "rps": float(r[o + 3]),
                "ir": float(r[o + 5]),
                "rps_rank": float(r[o + 4]),
                "ir_rank": float(r[o + 6]),
                "overall_rank": float(r[o + 2]),
                "position": int(r[o]),
            }
            if sheet == "Trial":
                row["evaluation"] = EVALUATIONS[0]
            elif sheet == "Month":
                row["evaluation"] = EVALUATIONS[int(r[8])]
            elif sheet == "Quarter":
                row["period"] = int(r[8])
            rows.append(row)
    wb.close()
    return pd.DataFrame(rows)

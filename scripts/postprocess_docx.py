"""Apply the UANL thesis layout to the docx rendered by Quarto.

Run automatically as a Quarto post-render hook (``thesis/_scripts/postrender.sh``)
or by hand:

    uv run python scripts/postprocess_docx.py thesis/_output/tesina.docx

What ``reference.docx`` cannot do on its own:

1. Front matter: cover and signature page copied from the official template
   (``thesis/templates/uanl_original.docx``, including the logo) and filled from the
   ``uanl`` key of ``thesis/_quarto.yml``; dedication; table of contents, list of
   tables and list of figures as Word fields; acknowledgements.
2. Chapter headings: "CAPÍTULO N" label above each numbered level-1 heading, and a
   hidden TC field so the table of contents shows "N. TÍTULO" as in the template.
3. Captions: Quarto emits paragraphs with two ``w:pPr`` (invalid OOXML) and uses the
   same style for figure and table captions; they are normalized so the lists of
   figures and tables can be built from their styles.
4. Page numbering: front matter and body are separate sections, each numbered from 1;
   the cover has no page number.
5. Word is asked to update fields on open, which fills in the indices.
"""

from __future__ import annotations

import argparse
import copy
import io
import os
import re
import sys
from pathlib import Path

import yaml
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

REPO_ROOT = Path(__file__).resolve().parent.parent
THESIS_DIR = REPO_ROOT / "thesis"
DEFAULT_CONFIG = THESIS_DIR / "_quarto.yml"
DEFAULT_TEMPLATE = THESIS_DIR / "templates" / "uanl_original.docx"
THESIS_OUTPUT_NAME = "tesina.docx"

CHAPTER_LABEL_STYLE = "UANLChapterLabel"
FRONT_HEADING_STYLE = "UANLFrontHeading"
INDEX_HEADER_STYLE = "UANLIndexHeader"
FIELD_PLACEHOLDER = "Actualice los campos para generar este índice (clic derecho → Actualizar)."

# Upper-case Arial 13 pt fits ~48 characters in the 15.1 cm text width of the cover.
COVER_CHARS_PER_LINE = 48

NUMBERED_HEADING_RE = re.compile(r"^(\d+)\.?\s+(.*)$")
R_EMBED = qn("r:embed")

# Elements that must follow w:updateFields in CT_Settings (schema order).
SETTINGS_AFTER_UPDATE_FIELDS = [
    "hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr",
    "attachedSchema", "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats",
    "doNotAutoCompressPictures", "forceUpgrade", "captions", "readModeInkLockDown",
    "smartTagType", "schemaLibrary", "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol",
    "listSeparator",
]  # fmt: skip


# --------------------------------------------------------------------------- XML helpers


def _el(tag: str, **attrs: str) -> OxmlElement:
    el = OxmlElement(tag)
    for key, value in attrs.items():
        el.set(qn(f"w:{key}"), str(value))
    return el


def _style_id(p) -> str | None:
    ps = p.find(f"{qn('w:pPr')}/{qn('w:pStyle')}")
    return ps.get(qn("w:val")) if ps is not None else None


def _text(el) -> str:
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def _paragraph(style: str | None = None, text: str | None = None, **ppr_children) -> OxmlElement:
    p = _el("w:p")
    ppr = _el("w:pPr")
    if style:
        ppr.append(_el("w:pStyle", val=style))
    for child in ppr_children.values():
        ppr.append(child)
    p.append(ppr)
    if text is not None:
        p.append(_run(text))
    return p


def _run(text: str, rpr: OxmlElement | None = None) -> OxmlElement:
    r = _el("w:r")
    if rpr is not None:
        r.append(copy.deepcopy(rpr))
    parts = text.split("\t")
    for i, part in enumerate(parts):
        if i:
            r.append(_el("w:tab"))
        if part:
            t = _el("w:t")
            t.text = part
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            r.append(t)
    return r


def _field_runs(instr: str, placeholder: str | None, hidden: bool = False) -> list:
    def fld(kind: str) -> OxmlElement:
        r = _el("w:r")
        if hidden:
            rpr = _el("w:rPr")
            rpr.append(_el("w:vanish"))
            r.append(rpr)
        r.append(_el("w:fldChar", fldCharType=kind))
        return r

    instr_run = _el("w:r")
    if hidden:
        rpr = _el("w:rPr")
        rpr.append(_el("w:vanish"))
        instr_run.append(rpr)
    instr_el = _el("w:instrText")
    instr_el.text = f" {instr} "
    instr_el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    instr_run.append(instr_el)

    runs = [fld("begin"), instr_run]
    if placeholder is not None:
        runs += [fld("separate"), _run(placeholder)]
    runs.append(fld("end"))
    return runs


def set_text(p, text: str) -> None:
    """Replace the paragraph's runs with one run holding ``text`` (keeps the first run's format)."""
    runs = p.findall(qn("w:r"))
    rpr = runs[0].find(qn("w:rPr")) if runs else None
    rpr = copy.deepcopy(rpr) if rpr is not None else None
    for r in runs:
        p.remove(r)
    p.append(_run(text, rpr))


# ------------------------------------------------------------------------------ config


def load_config(path: Path) -> dict:
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    cfg = data.get("uanl") or {}
    crossref = data.get("crossref") or {}
    cfg.setdefault("tbl-title", crossref.get("tbl-title", "Tabla"))
    cfg.setdefault("fig-title", crossref.get("fig-title", "Figura"))
    return cfg


def _value(cfg: dict, key: str, placeholder: str) -> str:
    value = cfg.get(key)
    return str(value).strip() if value else placeholder


# ---------------------------------------------------------------------------- captions


def normalize_captions(body, tbl_title: str) -> int:
    """Merge duplicated pPr and give table captions their own style. Returns #captions."""
    count = 0
    for p in body.iter(qn("w:p")):
        pprs = p.findall(qn("w:pPr"))
        style = next(
            (ps.get(qn("w:val")) for ppr in pprs if (ps := ppr.find(qn("w:pStyle"))) is not None),
            None,
        )
        if style not in ("ImageCaption", "TableCaption", "Caption"):
            continue
        if style != "Caption" and _text(p).startswith(f"{tbl_title} "):
            style = "TableCaption"
        for ppr in pprs:
            p.remove(ppr)
        ppr = _el("w:pPr")
        ppr.append(_el("w:pStyle", val=style))
        p.insert(0, ppr)
        count += 1
    return count


# ---------------------------------------------------------------------------- chapters


def format_chapters(body) -> int:
    """Insert the "CAPÍTULO N" label and a TC field for every level-1 heading."""
    chapters = 0
    for p in list(body.iterchildren(qn("w:p"))):
        if _style_id(p) != "Heading1":
            continue
        text = _text(p).strip()
        match = NUMBERED_HEADING_RE.match(text)
        if match:
            number, title = match.groups()
            set_text(p, title)
            toc_entry = f"{number}. {title.upper()}"
            label = _paragraph(CHAPTER_LABEL_STYLE, f"Capítulo {number}")
            p.addprevious(label)
            # The label carries the page break and top space; the heading follows it.
            ppr = p.find(qn("w:pPr"))
            ppr.append(_el("w:pageBreakBefore", val="0"))
            ppr.append(_el("w:spacing", before="0"))
            _sort_ppr(ppr)
            chapters += 1
        else:
            toc_entry = text
        for run in _field_runs(f'TC "{toc_entry}" \\l 1', None, hidden=True):
            p.append(run)
    return chapters


PPR_ORDER = [
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "pBdr", "tabs", "spacing", "ind",
    "jc", "outlineLvl", "rPr", "sectPr",
]  # fmt: skip
SECT_PR_ORDER = [
    "headerReference", "footerReference", "footnotePr", "endnotePr", "type", "pgSz", "pgMar",
    "paperSrc", "pgBorders", "lnNumType", "pgNumType", "cols", "formProt", "vAlign",
    "noEndnote", "titlePg", "textDirection", "bidi", "rtlGutter", "docGrid",
]  # fmt: skip


def _sort_children(parent, order: list[str]) -> None:
    """Reorder children to the OOXML schema sequence (stable for unknown/equal tags)."""
    rank = {qn(f"w:{name}"): i for i, name in enumerate(order)}
    for child in sorted(parent, key=lambda c: rank.get(c.tag, len(order))):
        parent.append(child)


def _sort_ppr(ppr) -> None:
    _sort_children(ppr, PPR_ORDER)


# ------------------------------------------------------------------------- front matter


def _find_index(elements, predicate, start: int = 0) -> int:
    for i in range(start, len(elements)):
        if predicate(elements[i]):
            return i
    raise ValueError("template layout not recognized")


def _copy_images(elements, template_doc, doc) -> None:
    """Re-link images in elements copied from the template into ``doc``."""
    for el in elements:
        for blip in el.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip"):
            image_part = template_doc.part.related_parts[blip.get(R_EMBED)]
            rid, _ = doc.part.get_or_add_image(io.BytesIO(image_part.blob))
            blip.set(R_EMBED, rid)
        for i, doc_pr in enumerate(
            el.iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr")
        ):
            doc_pr.set("id", str(9000 + i))


def _wrapped_lines(text: str, width: int = COVER_CHARS_PER_LINE) -> int:
    """Greedy word-wrap estimate of how many lines ``text`` takes on the cover."""
    lines, current = 1, 0
    for word in text.split():
        needed = len(word) if current == 0 else current + 1 + len(word)
        if needed > width and current:
            lines, current = lines + 1, len(word)
        else:
            current = needed
    return lines


def _fit_cover(cover: list, extra_lines: int) -> list:
    """Drop one empty paragraph per extra title line so the cover stays on one page.

    The template cover fills the page with a one-line title. Empty paragraphs are taken
    from the largest gaps first (lowest first among equals); every gap keeps at least one.
    """
    gaps: list[list] = []
    previous_blank = False
    for el in cover:
        blank = not _text(el).strip() and next(el.iter(qn("w:drawing")), None) is None
        if blank:
            if not previous_blank:
                gaps.append([])
            gaps[-1].append(el)
        previous_blank = blank
    removed = set()
    for _ in range(max(extra_lines, 0)):
        # Largest gap; among equals, the lowest on the page.
        gap = max(reversed(gaps), key=len)
        if len(gap) <= 1:
            break
        removed.add(id(gap.pop()))
    return [el for el in cover if id(el) not in removed]


def cover_and_signatures(template_doc, doc, cfg: dict) -> list:
    """Cover and signature page from the official template, with placeholders filled."""
    elements = list(template_doc.element.body.iterchildren())
    cover_end = _find_index(elements, lambda e: _text(e).strip().startswith("(Mes)"))
    sig_start = _find_index(
        elements, lambda e: _text(e).strip() == "Universidad Autónoma de Nuevo León", cover_end
    )
    sig_end = _find_index(
        elements, lambda e: _text(e).strip().startswith("San Nicolás de los Garza"), sig_start
    )
    cover = [copy.deepcopy(e) for e in elements[: cover_end + 1]]
    signatures = [copy.deepcopy(e) for e in elements[sig_start : sig_end + 1]]

    title = _value(cfg, "title", "")
    cover = _fit_cover(cover, _wrapped_lines(title.upper()) - 1)
    author = _value(cfg, "author", "(nombre completo del alumno, tipo oración, sin abreviaciones)")
    degree = _value(cfg, "degree", "Maestría en Ciencia de Datos")
    year = _value(cfg, "year", "(Año)")
    month = _value(cfg, "month", "(Mes)")
    reviewers = list(cfg.get("reviewers") or []) + ["", ""]
    reviewer_names = [str(r).strip() or "(Título y nombre)" for r in reviewers[:2]]

    for p in cover:
        text = _text(p).strip()
        if text == "TITULO":
            set_text(p, title.upper() if title else "TITULO")
        elif text == "MARCO ANTONIO OBREGÓN FLORES" and cfg.get("author"):
            set_text(p, author.upper())
        elif text == "MAESTRÍA EN CIENCIA DE DATOS":
            set_text(p, degree.upper())
        elif text.startswith("(Mes)"):
            set_text(p, f"{month}, {year}")

    director_done = False
    for p in signatures:
        text = _text(p).strip()
        if text.startswith("Los miembros del Comité"):
            set_text(
                p,
                "Los miembros del Comité de Tesis recomendamos que la Tesina "
                f"“{title or '(título del trabajo de investigación)'}”, realizada por "
                f"{_value(cfg, 'student', '(el, la) alumno(a)')} {author}, con número de "
                f"matrícula {_value(cfg, 'student-id', '(matrícula)')}, sea aceptada para su "
                f"defensa como opción al grado de {degree}.",
            )
        elif text == "(Título y nombre)" and not director_done:
            set_text(p, _value(cfg, "director", "(Título y nombre)"))
            director_done = True
        elif re.fullmatch(r"_+\s+_+", text):
            _two_columns(p, "_" * 30, "_" * 30)
        elif re.fullmatch(r"\(Título y nombre\)\s+\(Título y nombre\)", text):
            _two_columns(p, *reviewer_names)
        elif re.fullmatch(r"Revisor\s+Revisor", text):
            _two_columns(p, "Revisor", "Revisor")
        elif text.startswith("Dra. Azucena") and cfg.get("coordinator"):
            set_text(p, str(cfg["coordinator"]))
        elif text.startswith("Coordinadora") and cfg.get("coordinator-role"):
            set_text(p, str(cfg["coordinator-role"]))

    # The signature page starts on its own page.
    first_ppr = signatures[0].find(qn("w:pPr"))
    if first_ppr is None:
        first_ppr = _el("w:pPr")
        signatures[0].insert(0, first_ppr)
    first_ppr.insert(0, _el("w:pageBreakBefore"))
    _sort_ppr(first_ppr)

    elements = cover + signatures
    for el in elements:
        for tag in ("w:bookmarkStart", "w:bookmarkEnd"):
            for bm in el.iter(qn(tag)):
                bm.getparent().remove(bm)
    _copy_images(elements, template_doc, doc)
    return elements


def _two_columns(p, left: str, right: str) -> None:
    """Rewrite a space-aligned two-column line with center tabs at 1/4 and 3/4 width."""
    ppr = p.find(qn("w:pPr"))
    if ppr is None:
        ppr = _el("w:pPr")
        p.insert(0, ppr)
    for tag in ("w:jc", "w:tabs", "w:ind"):
        for old in ppr.findall(qn(tag)):
            ppr.remove(old)
    width = 12240 - 2268 - 1418
    tabs = _el("w:tabs")
    tabs.append(_el("w:tab", val="center", pos=str(width // 4)))
    tabs.append(_el("w:tab", val="center", pos=str(3 * width // 4)))
    ppr.append(tabs)
    ppr.append(_el("w:jc", val="left"))
    _sort_ppr(ppr)
    set_text(p, f"\t{left}\t{right}")


def _index(title: str, header: str, instr: str) -> list:
    p = _paragraph()
    for run in _field_runs(instr, FIELD_PLACEHOLDER):
        p.append(run)
    return [
        _paragraph(FRONT_HEADING_STYLE, title),
        _paragraph(INDEX_HEADER_STYLE, header),
        p,
    ]


def preliminaries(cfg: dict) -> list:
    """Dedication, indices and acknowledgements. A value of ``false`` omits a section."""
    elements = []
    dedication = cfg.get("dedication")
    if dedication is not False:
        elements.append(_paragraph(FRONT_HEADING_STYLE, "Dedicatoria"))
        elements += [_paragraph() for _ in range(3)]
        p = _paragraph(jc=_el("w:jc", val="right"))
        rpr = _el("w:rPr")
        rpr.append(_el("w:i"))
        p.append(
            _run(
                str(dedication).strip()
                or "(Dedicatoria opcional justificada a la derecha en itálicas)",
                rpr,
            )
        )
        elements.append(p)

    elements += _index("Índice de contenido", "Capítulo\tPágina", 'TOC \\o "2-3" \\h \\z \\f')
    elements += _index(
        "Índice de tablas", f"{cfg['tbl-title']}.\tPágina", 'TOC \\h \\z \\t "Table Caption,1"'
    )
    elements += _index(
        "Índice de figuras", f"{cfg['fig-title']}.\tPágina", 'TOC \\h \\z \\t "Image Caption,1"'
    )

    acknowledgements = cfg.get("acknowledgements")
    if acknowledgements is not False:
        elements.append(_paragraph(FRONT_HEADING_STYLE, "Agradecimientos"))
        text = str(acknowledgements).strip() or (
            "(Aquí se escriben los agradecimientos a personas o instituciones que apoyaron "
            "la elaboración de este trabajo de investigación. Este apartado es opcional.)"
        )
        for para in text.split("\n\n"):
            elements.append(_paragraph("BodyText", para.strip()))
    return elements


# ---------------------------------------------------------------------------- sections


def split_sections(body, last_front_matter) -> None:
    """Front matter and body as separate sections, each numbered from 1; no number on cover."""
    final_sect = body.find(qn("w:sectPr"))
    for old in final_sect.findall(qn("w:titlePg")) + final_sect.findall(qn("w:pgNumType")):
        final_sect.remove(old)
    front_sect = copy.deepcopy(final_sect)
    for sect, title_page in ((front_sect, True), (final_sect, False)):
        sect.append(_el("w:pgNumType", start="1"))
        if title_page:
            sect.append(_el("w:titlePg"))
        _sort_children(sect, SECT_PR_ORDER)
    ppr = last_front_matter.find(qn("w:pPr"))
    if ppr is None:
        ppr = _el("w:pPr")
        last_front_matter.insert(0, ppr)
    ppr.append(front_sect)
    _sort_ppr(ppr)


def request_field_update(doc) -> None:
    settings = doc.settings.element
    for old in settings.findall(qn("w:updateFields")):
        settings.remove(old)
    update = _el("w:updateFields", val="true")
    after = {qn(f"w:{name}") for name in SETTINGS_AFTER_UPDATE_FIELDS}
    for child in settings:
        if child.tag in after:
            child.addprevious(update)
            return
    settings.append(update)


# -------------------------------------------------------------------------------- main


def is_processed(body) -> bool:
    return any(_style_id(p) == FRONT_HEADING_STYLE for p in body.iterchildren(qn("w:p")))


def process(docx_path: Path, config_path: Path, template_path: Path) -> None:
    cfg = load_config(config_path)
    doc = Document(docx_path)
    body = doc.element.body
    if is_processed(body):
        print(f"{docx_path}: already processed, skipping")
        return

    captions = normalize_captions(body, cfg["tbl-title"])
    chapters = format_chapters(body)

    template_doc = Document(template_path)
    front = cover_and_signatures(template_doc, doc, cfg) + preliminaries(cfg)
    for el in reversed(front):
        body.insert(0, el)
    split_sections(body, front[-1])
    request_field_update(doc)

    doc.core_properties.title = str(cfg.get("title") or "")
    doc.core_properties.author = str(cfg.get("author") or "")
    doc.save(docx_path)
    print(f"{docx_path}: {chapters} chapters, {captions} captions, front matter added")


def quarto_outputs() -> list[Path]:
    files = os.environ.get("QUARTO_PROJECT_OUTPUT_FILES", "")
    return [
        Path(f)
        for f in files.splitlines()
        if f.strip() and Path(f.strip()).name == THESIS_OUTPUT_NAME
    ]


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("docx", nargs="*", type=Path)
    parser.add_argument(
        "--from-quarto-env",
        action="store_true",
        help=f"process {THESIS_OUTPUT_NAME} from QUARTO_PROJECT_OUTPUT_FILES",
    )
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    args = parser.parse_args(argv)

    paths = list(args.docx)
    if args.from_quarto_env:
        paths += quarto_outputs()
    if not paths:
        print("no thesis docx to process", file=sys.stderr)
        return 0
    for path in paths:
        process(path.resolve(), args.config, args.template)
    return 0


if __name__ == "__main__":
    sys.exit(main())

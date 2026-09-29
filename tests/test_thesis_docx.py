"""Tests for scripts/build_reference_docx.py and scripts/postprocess_docx.py."""

import importlib.util
import sys
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "thesis" / "templates" / "uanl_original.docx"


def _load(name):
    spec = importlib.util.spec_from_file_location(name, ROOT / "scripts" / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


build_reference = _load("build_reference_docx")
postprocess = _load("postprocess_docx")

CONFIG = """\
crossref:
  tbl-title: Tabla
  fig-title: Figura
uanl:
  title: "Un título de prueba"
  author: "Nombre Apellido"
  student: "el alumno"
  student-id: "1234567"
  degree: "Maestría en Ciencia de Datos"
  month: "Enero"
  year: "2027"
  director: "Dr. Director"
  reviewers: ["Dra. Revisora", ""]
  dedication: false
  acknowledgements: "Gracias."
"""


@pytest.fixture(scope="module")
def reference(tmp_path_factory):
    path = tmp_path_factory.mktemp("ref") / "reference.docx"
    build_reference.build(TEMPLATE, path)
    return path


def _style_ids(doc):
    return {s.style_id for s in doc.styles}


def test_reference_defines_pandoc_and_custom_styles(reference):
    doc = Document(reference)
    expected = {
        "BodyText", "FirstParagraph", "Compact", "Heading1", "Heading2", "Heading3",
        "Heading4", "Heading5", "Bibliography", "ImageCaption", "TableCaption", "TOC1",
        "UANLChapterLabel", "UANLFrontHeading", "UANLIndexHeader",
    }  # fmt: skip
    assert expected <= _style_ids(doc)


def test_reference_keeps_template_page_setup_and_footer(reference):
    doc = Document(reference)
    assert len(doc.sections) == 1
    section = doc.sections[0]
    assert (section.page_width, section.page_height) == (12240 * 635, 15840 * 635)
    assert section.left_margin == 2268 * 635
    footer_xml = section.footer._element.xml
    assert "PAGE" in footer_xml


def _fake_pandoc_docx(reference, path):
    """Mimic Quarto's output: numbered headings and a caption with two pPr."""
    doc = Document(reference)
    body = doc.element.body
    for p in list(body.iterchildren(qn("w:p"))):
        body.remove(p)
    doc.add_paragraph("1. Introducción", style="heading 1")
    doc.add_paragraph("1.1 Antecedentes", style="heading 2")
    caption = doc.add_paragraph("Tabla 1. Resultados")
    caption.style = doc.styles["Image Caption"]
    extra = caption._p.makeelement(qn("w:pPr"), {})
    extra.append(caption._p.makeelement(qn("w:jc"), {qn("w:val"): "center"}))
    caption._p.insert(0, extra)
    doc.add_paragraph("Figura 1. Una figura", style="Image Caption")
    doc.add_paragraph("Referencias", style="heading 1")
    doc.save(path)


@pytest.fixture
def processed(reference, tmp_path):
    out = tmp_path / "tesina.docx"
    config = tmp_path / "_quarto.yml"
    config.write_text(CONFIG, encoding="utf-8")
    _fake_pandoc_docx(reference, out)
    postprocess.process(out, config, TEMPLATE)
    return Document(out)


def _paragraphs(doc):
    return [(p.style.style_id if p.style else None, p.text) for p in doc.paragraphs]


def test_chapter_label_and_tc_field(processed):
    paras = _paragraphs(processed)
    i = paras.index(("UANLChapterLabel", "Capítulo 1"))
    assert paras[i + 1] == ("Heading1", "Introducción")
    xml = processed.element.body.xml
    assert 'TC "1. INTRODUCCIÓN" \\l 1' in xml
    assert 'TC "Referencias" \\l 1' in xml
    assert ("Heading2", "1.1 Antecedentes") in paras


def test_captions_normalized(processed):
    paras = _paragraphs(processed)
    assert ("TableCaption", "Tabla 1. Resultados") in paras
    assert ("ImageCaption", "Figura 1. Una figura") in paras
    for p in processed.element.body.iter(qn("w:p")):
        assert len(p.findall(qn("w:pPr"))) <= 1


def test_front_matter_filled_from_config(processed):
    texts = [p.text for p in processed.paragraphs]
    assert "UN TÍTULO DE PRUEBA" in texts
    assert "NOMBRE APELLIDO" in texts
    assert "Enero, 2027" in texts
    assert "Dr. Director" in texts
    committee = next(t for t in texts if t.startswith("Los miembros del Comité"))
    assert "“Un título de prueba”" in committee
    assert "el alumno Nombre Apellido" in committee
    assert "matrícula 1234567" in committee
    assert "\tDra. Revisora\t(Título y nombre)" in texts
    assert "Dedicatoria" not in texts  # dedication: false omits it
    assert "Gracias." in texts
    assert ("UANLFrontHeading", "Índice de contenido") in _paragraphs(processed)


def test_indices_are_word_fields(processed):
    instr = " ".join(t.text for t in processed.element.body.iter(qn("w:instrText")))
    assert 'TOC \\o "2-3" \\h \\z \\f' in instr
    assert 'TOC \\h \\z \\t "Table Caption,1"' in instr
    assert 'TOC \\h \\z \\t "Image Caption,1"' in instr
    settings = processed.settings.element
    assert settings.find(qn("w:updateFields")).get(qn("w:val")) == "true"


def test_logo_copied(processed):
    blips = list(
        processed.element.body.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip")
    )
    assert blips
    rid = blips[0].get(qn("r:embed"))
    assert processed.part.related_parts[rid].content_type.startswith("image/")


def test_two_sections_numbered_from_one(processed):
    sections = processed.sections
    assert len(sections) == 2
    front, main = sections
    assert front.different_first_page_header_footer  # cover without page number
    assert not main.different_first_page_header_footer
    for section in sections:
        pg = section._sectPr.find(qn("w:pgNumType"))
        assert pg.get(qn("w:start")) == "1"


def test_idempotent(processed, tmp_path, capsys):
    out = tmp_path / "again.docx"
    processed.save(out)
    config = tmp_path / "_quarto.yml"
    postprocess.process(out, config, TEMPLATE)
    assert "already processed" in capsys.readouterr().out


def test_cover_fit_estimate():
    assert postprocess._wrapped_lines("TITULO") == 1
    assert postprocess._wrapped_lines("PALABRA " * 12) == 2
